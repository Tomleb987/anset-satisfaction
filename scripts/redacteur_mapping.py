#!/usr/bin/env python3
# =============================================================================
#  ANSET — Correspondance `req` → RÉDACTEUR, extraite des requêtes mensuelles.
#
#  POURQUOI CE SCRIPT EXISTE. Depuis le 08/09/2026 l'import attribue une réponse
#  au rédacteur (celui qui a établi le contrat et parlé au client) et non plus au
#  gestionnaire (qui suit le portefeuille). Pour les mois DÉJÀ diffusés, la
#  colonne « Redacteur » n'a jamais été lue : l'information n'existe nulle part en
#  base, et seuls les .xlsx source peuvent la rendre. Ce script les relit et
#  produit le SQL qui charge `public.import_redacteur`, puis appelle
#  `public.appliquer_redacteur()` (migration 20260908090000).
#
#  USAGE — passer les requêtes DANS L'ORDRE CHRONOLOGIQUE :
#     python3 scripts/redacteur_mapping.py "requete 2026-06.xlsx" "requete 2026-07.xlsx"
#     → écrit redacteur_import.sql (option -o) + un rapport à l'écran.
#  Le .sql se colle tel quel dans l'éditeur SQL Supabase (pas de `supabase login`
#  possible depuis cet environnement — voir README).
#
#  CE QU'IL NE FAUT PAS LUI DONNER : l'export « sinistres clos ». Il n'a pas de
#  colonne rédacteur, et son gestionnaire a réellement traité le dossier ; le
#  script le détecte et refuse. La migration écarte de toute façon tout `req`
#  portant un envoi `motif = 'sinistre'`.
#
#  LA CLÉ ET LE SLUG SONT CEUX DE L'APP, à la lettre : `req` = « Quittance », à
#  défaut « Dossier » ; slug = minuscules sans accents, tout ce qui n'est ni
#  lettre ni chiffre ni point devenant un point (`slugConseiller` de
#  satisfaction_anset.html). Une divergence d'un caractère créerait un conseiller
#  fantôme au lieu de rebasculer les réponses.
# =============================================================================

import argparse
import csv
import datetime as dt
import re
import sys
import unicodedata
from pathlib import Path

TAILLE_LOT = 500  # lignes par `insert` : l'éditeur SQL n'aime pas les requêtes géantes


def norm(v):
    """`norm` de l'app : texte, minuscules, sans accents, sans espaces de bord."""
    s = "" if v is None else str(v)
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.lower().strip()


def slug_conseiller(v):
    """`slugConseiller` de l'app."""
    s = re.sub(r"[^a-z0-9.]+", ".", norm(v))
    return s.strip(".")


def cellule(v):
    """Équivalent du `raw:false` de SheetJS : la valeur telle qu'affichée.

    openpyxl rend les entiers d'Excel en float (41633.0). Recopié tel quel, le
    `req` ne correspondrait à aucune ligne d'`envois_sondage`.
    """
    if v is None:
        return ""
    if isinstance(v, bool):
        return "OUI" if v else "NON"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def map_headers(entetes):
    """`mapReqHeaders` de l'app, réduit aux colonnes utiles ici.

    Même chaîne de priorités : « Redacteur » est testé AVANT « Gestionnaire »,
    et « sinistre » avant « quittance ».
    """
    idx = {}
    for i, brut in enumerate(entetes):
        n = norm(brut)
        if "sinistre" in n:
            idx.setdefault("sinistre", i)
        elif "quittance" in n:
            idx.setdefault("quittance", i)
        elif "dossier" in n:
            idx.setdefault("dossier", i)
        elif "redacteur" in n:
            idx.setdefault("redacteur", i)
        elif "gestionnaire" in n:
            idx.setdefault("gestionnaire", i)
    return idx


def lire_lignes(chemin):
    """Renvoie la liste de listes (première feuille), en-tête inclus."""
    suffixe = chemin.suffix.lower()
    if suffixe == ".xls":
        raise SystemExit(
            f"{chemin.name} : format .xls non lu ici. L'ouvrir et l'enregistrer en .xlsx."
        )
    if suffixe in (".csv", ".txt"):
        for encodage in ("utf-8-sig", "cp1252"):
            try:
                texte = chemin.read_text(encoding=encodage)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise SystemExit(f"{chemin.name} : encodage illisible.")
        dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=";,\t|")
        return [list(l) for l in csv.reader(texte.splitlines(), dialecte)]
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("openpyxl manquant : pip install openpyxl")
    classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    feuille = classeur[classeur.sheetnames[0]]
    lignes = [[cellule(c) for c in ligne] for ligne in feuille.iter_rows(values_only=True)]
    classeur.close()
    return lignes


def echappe(s):
    return s.replace("'", "''")


def main():
    ap = argparse.ArgumentParser(
        description="Extrait `req` → rédacteur des requêtes mensuelles et produit le SQL de rebascule."
    )
    ap.add_argument("fichiers", nargs="+", help="requêtes .xlsx / .csv, dans l'ordre chronologique")
    ap.add_argument("-o", "--sortie", default="redacteur_import.sql", help="fichier SQL à écrire")
    ap.add_argument(
        "--sans-appel",
        action="store_true",
        help="ne pas ajouter l'appel à appliquer_redacteur() (charge la table seulement)",
    )
    args = ap.parse_args()

    mapping = {}       # req -> redacteur (le dernier fichier gagne)
    origine = {}       # req -> nom du fichier qui a fourni la valeur retenue
    divergences = []   # (req, ancien, nouveau, fichier)
    diff_gest = 0      # req où rédacteur ≠ gestionnaire (le constat de la migration)
    sans_redacteur = 0
    sans_req = 0
    total = 0

    for nom in args.fichiers:
        chemin = Path(nom)
        if not chemin.exists():
            raise SystemExit(f"{nom} : fichier introuvable.")
        lignes = lire_lignes(chemin)
        if len(lignes) < 2:
            raise SystemExit(f"{chemin.name} : fichier vide.")
        idx = map_headers(lignes[0])
        if "redacteur" not in idx:
            raise SystemExit(
                f"{chemin.name} : pas de colonne « Redacteur ». "
                "Export « sinistres clos » ? Il reste attribué au gestionnaire, à raison."
            )
        if "quittance" not in idx and "dossier" not in idx:
            raise SystemExit(f"{chemin.name} : ni « Quittance » ni « Dossier » — pas de clé req.")

        lus = 0
        for ligne in lignes[1:]:
            def g(cle):
                i = idx.get(cle)
                if i is None or i >= len(ligne):
                    return ""
                return str(ligne[i] or "").strip()

            if not any(g(c) for c in ("quittance", "dossier", "redacteur", "gestionnaire")):
                continue  # ligne totalement vide (fin de feuille)
            total += 1
            lus += 1
            req = g("quittance") or g("dossier")
            redacteur = slug_conseiller(g("redacteur"))
            if not req:
                sans_req += 1
                continue
            if not redacteur:
                # Sans rédacteur, rien à rebasculer : le gestionnaire reste en place.
                sans_redacteur += 1
                continue
            if slug_conseiller(g("gestionnaire")) != redacteur:
                diff_gest += 1
            ancien = mapping.get(req)
            if ancien is not None and ancien != redacteur:
                divergences.append((req, ancien, redacteur, chemin.name))
            mapping[req] = redacteur
            origine[req] = chemin.name
        print(f"· {chemin.name} : {lus} lignes lues", file=sys.stderr)

    if not mapping:
        raise SystemExit("Aucune correspondance req → rédacteur : rien à écrire.")

    reqs = sorted(mapping)
    redacteurs = sorted(set(mapping.values()))
    horodatage = dt.datetime.now().strftime("%d/%m/%Y %H:%M")

    lignes_sql = [
        "-- =============================================================================",
        f"--  ANSET — Chargement de public.import_redacteur ({horodatage}).",
        "--  Généré par scripts/redacteur_mapping.py depuis :",
    ]
    lignes_sql += [f"--    · {Path(f).name}" for f in args.fichiers]
    lignes_sql += [
        f"--  {len(reqs)} correspondances, {len(redacteurs)} rédacteurs distincts.",
        "--",
        "--  À coller dans l'éditeur SQL Supabase. Une seule transaction : si l'appel à",
        "--  appliquer_redacteur() échoue, le chargement est annulé avec lui. Rejouable —",
        "--  un req déjà présent est mis à jour, la fonction ne réécrit que ce qui diffère.",
        "-- =============================================================================",
        "",
        "begin;",
        "",
    ]
    for depart in range(0, len(reqs), TAILLE_LOT):
        lot = reqs[depart : depart + TAILLE_LOT]
        lignes_sql.append("insert into public.import_redacteur (req, redacteur) values")
        lignes_sql += [
            f"  ('{echappe(r)}', '{echappe(mapping[r])}'){',' if i < len(lot) - 1 else ''}"
            for i, r in enumerate(lot)
        ]
        lignes_sql.append(
            "on conflict (req) do update set redacteur = excluded.redacteur, charge_le = now();"
        )
        lignes_sql.append("")
    if not args.sans_appel:
        lignes_sql += [
            "-- Rebascule envois / réponses / leads. Les campagnes sinistre sont écartées",
            "-- par la fonction elle-même (reqs_ignores_sin dans le retour).",
            "select * from public.appliquer_redacteur();",
            "",
        ]
    lignes_sql += ["commit;", ""]

    Path(args.sortie).write_text("\n".join(lignes_sql), encoding="utf-8")

    print(f"\n{args.sortie} écrit.")
    print(f"  lignes lues .................. {total}")
    print(f"  correspondances req → rédacteur {len(reqs)}")
    print(f"  rédacteurs distincts ......... {len(redacteurs)}")
    print(f"  rédacteur ≠ gestionnaire ..... {diff_gest}" + (f" ({100*diff_gest/total:.1f} %)" if total else ""))
    print(f"  sans rédacteur (non touchées)  {sans_redacteur}")
    if sans_req:
        print(f"  sans req (ignorées) .......... {sans_req}")
    if divergences:
        print(f"\n  {len(divergences)} req changent de rédacteur d'un mois à l'autre.")
        print("  Le DERNIER fichier de la ligne de commande gagne — d'où l'ordre chronologique.")
        for req, ancien, nouveau, fichier in divergences[:10]:
            print(f"    {req} : {ancien} → {nouveau} ({fichier})")
        if len(divergences) > 10:
            print(f"    … et {len(divergences)-10} autres.")


if __name__ == "__main__":
    main()
